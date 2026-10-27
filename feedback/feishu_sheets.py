"""Merge completed detail exports into a new tab in a configured Feishu Sheet."""
from datetime import date, datetime
import hashlib
import json
import re
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel

from .exports.excel import COLUMNS
from .feishu_api import ApiError, Client
from .storage import atomic_json

# Feishu's API accepts predefined formats, not every Excel custom format.
# Verified against the API; hyphenated date formats return code 90204.
DATE_FORMATTER = "yyyy/MM/dd HH:mm:ss"


def settings_for(config):
    settings = config.get("feishu", {}).get("spreadsheet", {})
    return settings if settings.get("enabled") else None


def destination(settings):
    url = settings.get("url", "")
    parsed = urlsplit(url)
    if parsed.scheme != "https" or not re.fullmatch(r"[A-Za-z0-9-]+\.feishu\.cn", parsed.netloc):
        raise ValueError("飞书表格配置 url 必须是有效的 https://企业.feishu.cn 链接。")
    match = re.fullmatch(r"/(wiki|sheets)/([A-Za-z0-9_-]+)/?", parsed.path)
    if not match:
        raise ValueError("飞书目标必须是电子表格链接（/sheets/）或其知识库链接（/wiki/）。")
    return match.groups()


def resolve(client, settings):
    kind, identifier = destination(settings)
    explicit = settings.get("spreadsheet_token")
    if explicit:
        if not re.fullmatch(r"[A-Za-z0-9_-]+", explicit):
            raise ValueError("飞书 spreadsheet_token 配置无效。")
        return explicit
    if kind == "sheets":
        return identifier
    try:
        node = client.call("GET", "/wiki/v2/spaces/get_node?" + urlencode({"token": identifier}))["node"]
    except ApiError as error:
        raise ValueError(f"无法解析 Wiki 中的电子表格。请为应用开通知识库节点读取权限并发布，或配置已核实的 spreadsheet_token。{error}") from None
    if node.get("obj_type") != "sheet" or not re.fullmatch(r"[A-Za-z0-9_-]+", node.get("obj_token", "")):
        raise ValueError("目标 Wiki 节点不是飞书电子表格，请核对链接。")
    return node["obj_token"]


def detail_rows(folder):
    """Read only the app's export; retain period membership and literal IDs."""
    book = load_workbook(folder / "feedback_details.xlsx", read_only=True, data_only=False)
    try:
        headers = ["统计期间", *COLUMNS.values()]
        rows = [headers]
        for title, period in (("本期反馈明细", "本期"), ("上周同期反馈明细", "上周同期")):
            if title not in book.sheetnames:
                raise ValueError("报告 Excel 缺少预期的明细工作表。")
            sheet = book[title]
            values = list(sheet.iter_rows(min_row=4, values_only=True))
            if not values or list(values[0]) != list(COLUMNS.values()):
                raise ValueError("报告 Excel 明细字段不匹配，无法推送。")
            for record in values[1:]:
                if all(value is None for value in record):
                    continue
                converted = []
                for label, value in zip(COLUMNS.values(), record):
                    if value is None:
                        value = ""
                    if label == "反馈时间" and isinstance(value, (datetime, date)):
                        value = to_excel(value)
                    elif isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                        # Literal user feedback must never be interpreted as a formula.
                        value = "'" + value
                    if isinstance(value, str) and len(value) > 50000:
                        raise ValueError("明细文字超过飞书单元格长度限制，未截断，请检查数据。")
                    converted.append(value)
                rows.append([period, *converted])
        return rows
    finally:
        book.close()


def sheet_url(url, sheet_id):
    parsed = urlsplit(url)
    query = [(key, value) for key, value in parse_qsl(parsed.query) if key not in {"sheet", "from"}]
    return urlunsplit(parsed._replace(query=urlencode([*query, ("sheet", sheet_id)]), fragment=""))


def saved_state(folder):
    path = folder / "feishu_sheet.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}


def sync(folder, config):
    settings = settings_for(config)
    if not settings:
        return None
    destination(settings)
    batch_rows = settings.get("batch_rows", 500)
    if type(batch_rows) is not int or not 1 <= batch_rows <= 5000:
        raise ValueError("飞书 batch_rows 必须是 1 到 5000 的整数。")
    rows = detail_rows(folder)
    client = Client(settings)
    path = folder / "feishu_sheet.json"
    saved = saved_state(folder)
    target = settings["url"]
    fingerprint = hashlib.sha256(json.dumps(rows, ensure_ascii=False).encode()).hexdigest()
    if saved and (saved["target"] != target or saved["fingerprint"] != fingerprint):
        raise ValueError("此报告的飞书目标或明细已经变化。请恢复原配置，或生成新报告后推送。")
    if not saved:
        # Periods are from the completed report, never mutable form controls.
        first = load_workbook(folder / "feedback_details.xlsx", read_only=True)
        try:
            period_text = first["本期反馈明细"]["A2"].value or ""
        finally:
            first.close()
        dates = re.findall(r"\d{4}-\d{2}-\d{2}", period_text)
        if len(dates) != 2:
            raise ValueError("报告明细缺少日期范围。")
        title = f"{dates[0].replace('-', '.')}-{dates[1].replace('-', '.')}"
        saved = {"target": target, "fingerprint": fingerprint, "title": title,
                 "written_rows": 0, "total_rows": len(rows), "status": "pending"}
        atomic_json(path, saved)
    if saved.get("status") == "complete":
        return saved

    def checkpoint(**values):
        saved.update(values)
        atomic_json(path, saved)

    try:
        token = saved.get("spreadsheet_token") or resolve(client, settings)
        checkpoint(spreadsheet_token=token, status="syncing", error="")
        root = f"/sheets/v2/spreadsheets/{token}"
        sheets = client.call("GET", f"/sheets/v3/spreadsheets/{token}/sheets/query").get("sheets", [])
        if not saved.get("sheet_id"):
            # Record the date-only title after checking for collisions. A
            # creation timeout can then resume the original pending report.
            if not saved.get("create_title"):
                titles = {item["title"] for item in sheets}
                create_title = saved["title"]
                if create_title in titles:
                    raise ValueError(f"飞书已存在页签「{create_title}」。请先重命名原页签后再推送新报告，以免覆盖已有明细。")
                checkpoint(create_title=create_title)
            existing = [item for item in sheets if item["title"] == saved["create_title"]]
            if existing:
                sheet_id = existing[0]["sheet_id"]
            else:
                result = client.call("POST", root + "/sheets_batch_update", {"requests": [{"addSheet": {
                    "properties": {"title": saved["create_title"], "index": len(sheets)}}}]})
                sheet_id = result["replies"][0]["addSheet"]["properties"]["sheetId"]
            checkpoint(sheet_id=sheet_id, url=sheet_url(target, sheet_id))
            sheets = client.call("GET", f"/sheets/v3/spreadsheets/{token}/sheets/query").get("sheets", [])
        sheet_id = saved["sheet_id"]
        if not re.fullmatch(r"[A-Za-z0-9_-]+", sheet_id):
            raise ValueError("飞书返回的页签 ID 无效。")
        matching = [item for item in sheets if item["sheet_id"] == sheet_id]
        if not matching:
            raise ValueError("已创建的报告页签不存在，请检查是否被手动删除。")
        end_column = get_column_letter(len(rows[0]))
        if not saved.get("prepared"):
            # Only enlarge our own new tab, querying dimensions again on retry.
            grid = matching[0].get("grid_properties", {})
            for dimension, key, required in (("ROWS", "row_count", max(2, len(rows))),
                                               ("COLUMNS", "column_count", len(rows[0]))):
                count = grid.get(key)
                if type(count) is not int:
                    raise ValueError("无法读取新页签的行列容量，请重试。")
                while count < required:
                    length = min(5000, required - count)
                    client.call("POST", root + "/dimension_range", {"dimension": {
                        "sheetId": sheet_id, "majorDimension": dimension, "length": length}})
                    count += length
            client.call("POST", root + "/sheets_batch_update", {"requests": [{"updateSheet": {
                "properties": {"sheetId": sheet_id, "frozenRowCount": 1}}}]})
            client.call("PUT", root + "/styles_batch_update", {"data": [
                {"ranges": [f"{sheet_id}!A1:{end_column}1"], "style": {
                    "font": {"bold": True}, "backColor": "#2563EB", "foreColor": "#FFFFFF"}}]})
            checkpoint(prepared=True)
        offset = saved["written_rows"]
        while offset < len(rows):
            batch = rows[offset:offset + batch_rows]
            # Keep requests small even with long Chinese/Arabic feedback.
            while len(json.dumps(batch, ensure_ascii=False).encode()) > 1500000 and len(batch) > 1:
                batch = batch[:max(1, len(batch) // 2)]
            date_column = get_column_letter(rows[0].index("反馈时间") + 1)
            styles = [{"ranges": [f"{sheet_id}!A{offset + 1}:{end_column}{offset + len(batch)}"],
                       "style": {"formatter": "@", "vAlign": 0}}]
            if offset + len(batch) > 1:
                styles.append({"ranges": [f"{sheet_id}!{date_column}{max(2, offset + 1)}:{date_column}{offset + len(batch)}"],
                               "style": {"formatter": DATE_FORMATTER}})

            client.call("PUT", root + "/styles_batch_update", {"data": styles})
            result = client.call("PUT", root + "/values", {"valueRange": {
                "range": f"{sheet_id}!A{offset + 1}:{end_column}{offset + len(batch)}", "values": batch}})
            if result.get("updatedRows") != len(batch) or result.get("updatedColumns") != len(rows[0]):
                raise ValueError("飞书返回的写入行列数不匹配，请再次推送以继续。")
            offset += len(batch)
            checkpoint(written_rows=offset)
        checkpoint(status="complete")
        return saved
    except (ValueError, KeyError, TypeError) as error:
        message = str(error) if isinstance(error, ValueError) else "飞书接口响应格式不符合预期，请核对配置后重试。"
        checkpoint(status="failed", error=message)
        raise ValueError(message) from None


def main():
    """Read-only permission check; does not create sheets or send messages."""
    import argparse
    from dotenv import load_dotenv
    from .config import ROOT, load_config
    parser = argparse.ArgumentParser(description="检查游戏飞书表格连接（不会写入或发送消息）")
    parser.add_argument("--check", metavar="GAME", required=True)
    args = parser.parse_args()
    load_dotenv(ROOT / ".env")
    try:
        settings = settings_for(load_config(args.check))
        if not settings:
            raise ValueError("此游戏未启用飞书表格同步。")
        client = Client(settings)
        token = resolve(client, settings)
        metadata = client.call("GET", f"/sheets/v3/spreadsheets/{token}")
        sheets = client.call("GET", f"/sheets/v3/spreadsheets/{token}/sheets/query").get("sheets", [])
        permission = metadata.get("spreadsheet", {}).get("user_permission", "未知")
        print(f"连接成功：已有 {len(sheets)} 个页签；访问权限：{permission}。未写入数据或发送消息。")
    except ValueError as error:
        parser.exit(1, str(error) + "\n")


if __name__ == "__main__":
    main()
