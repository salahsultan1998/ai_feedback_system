from copy import deepcopy
import json
import os
from unittest.mock import patch

from feedback import feishu, feishu_sheets
from feedback.feishu_api import ApiError, Client, _tokens
from feedback.workflow import run_report
from tests.helpers import BaseTest, fake_request


class FakeSheets:
    def __init__(self):
        self.calls = []
        self.sheets = [{"sheet_id": "history", "title": "历史明细", "grid_properties": {
            "row_count": 200, "column_count": 26}}]
        self.fail_create = False
        self.fail_range = None
        self.writes = {}

    def call(self, method, path, payload=None):
        self.calls.append((method, path, deepcopy(payload)))
        if "/wiki/" in path:
            return {"node": {"obj_type": "sheet", "obj_token": "target"}}
        if path.endswith("/sheets/query"):
            return {"sheets": deepcopy(self.sheets)}
        if path.endswith("/sheets_batch_update"):
            request = payload["requests"][0]
            if "addSheet" in request:
                props = request["addSheet"]["properties"]
                self.sheets.append({"sheet_id": "tab1", "title": props["title"],
                                    "grid_properties": {"row_count": 200, "column_count": 26}})
                if self.fail_create:
                    self.fail_create = False
                    raise ApiError("创建请求超时")
                return {"replies": [{"addSheet": {"properties": {"sheetId": "tab1"}}}]}
            return {}
        if path.endswith("/styles_batch_update"):
            for item in payload["data"]:
                formatter = item["style"].get("formatter")
                if formatter is not None and formatter not in {"@", "yyyy/MM/dd HH:mm:ss"}:
                    raise ApiError("invalid formatter", 90204)
            return {}
        if path.endswith("/dimension_range"):
            dimension = payload["dimension"]
            key = "row_count" if dimension["majorDimension"] == "ROWS" else "column_count"
            self.sheets[-1]["grid_properties"][key] += dimension["length"]
            return {"addCount": dimension["length"]}
        if path.endswith("/values"):
            value = payload["valueRange"]
            if value["range"] == self.fail_range:
                self.fail_range = None
                raise ApiError("写入请求超时")
            self.writes[value["range"]] = value["values"]
            return {"updatedRows": len(value["values"]), "updatedColumns": len(value["values"][0])}
        raise AssertionError((method, path, payload))


class SheetTests(BaseTest):
    def setUp(self):
        super().setUp()
        with patch("feedback.ai_client.request_json", side_effect=fake_request):
            run_report(self.root, self.config, "2026-10-19", "2026-10-22", False,
                       "report", lambda **v: None)
        self.folder = self.root / "reports/report"
        self.config["feishu"]["spreadsheet"] = {"enabled": True,
            "url": "https://example.feishu.cn/wiki/node", "batch_rows": 2}
        self.server = FakeSheets()
        mock = patch("feedback.feishu_sheets.Client", return_value=self.server)
        mock.start()
        self.addCleanup(mock.stop)

    def creates(self):
        return [payload for _, path, payload in self.server.calls
                if path.endswith("/sheets_batch_update") and "addSheet" in payload["requests"][0]]

    def test_merge_fields_periods_ids_dates_and_literal_feedback(self):
        rows = feishu_sheets.detail_rows(self.folder)
        self.assertEqual(len(rows), 5)
        self.assertEqual(len(rows[0]), 11)
        self.assertNotIn("原标签", rows[0])
        self.assertNotIn("用户主动选择的标签", rows[0])
        self.assertEqual(rows[0][-1], "是否建议")
        self.assertEqual(rows[1][-1], "是")
        self.assertEqual([row[0] for row in rows[1:]], ["本期", "本期", "上周同期", "上周同期"])
        self.assertEqual(rows[1][3], "004")
        self.assertEqual(rows[1][4], "0009")
        self.assertIsInstance(rows[1][6], (int, float))
        self.assertEqual(rows[1][8], "'=1+1")

    def test_sync_and_repeat_preserve_existing_tabs(self):
        result = feishu_sheets.sync(self.folder, self.config)
        self.assertEqual(result["status"], "complete")
        self.assertEqual(result["written_rows"], 5)
        self.assertIn("sheet=tab1", result["url"])
        self.assertEqual(self.server.sheets[-1]["title"], "2026.10.19-2026.10.22")
        self.assertEqual(len(self.creates()), 1)
        count = len(self.server.calls)
        self.assertEqual(feishu_sheets.sync(self.folder, self.config), result)
        self.assertEqual(len(self.server.calls), count)
        self.assertEqual(self.server.sheets[0]["title"], "历史明细")
        self.assertEqual(sum(map(len, self.server.writes.values())), 5)
        self.assertTrue(all(key.startswith("tab1!") for key in self.server.writes))

    def test_creation_timeout_recovers_same_tab(self):
        self.server.fail_create = True
        with self.assertRaisesRegex(ValueError, "创建请求超时"):
            feishu_sheets.sync(self.folder, self.config)
        result = feishu_sheets.sync(self.folder, self.config)
        self.assertEqual(result["status"], "complete")
        self.assertEqual(len(self.creates()), 1)
        self.assertEqual(len(self.server.sheets), 2)

    def test_existing_date_title_does_not_overwrite_another_report(self):
        self.server.sheets[0]["title"] = "2026.10.19-2026.10.22"
        for _ in range(2):
            with self.assertRaisesRegex(ValueError, "请先重命名原页签"):
                feishu_sheets.sync(self.folder, self.config)
        self.assertFalse(self.creates())
        self.assertFalse(self.server.writes)
        self.assertNotIn("create_title", feishu_sheets.saved_state(self.folder))

    def test_legacy_pending_tab_is_resumed_with_its_original_name(self):
        self.server.fail_create = True
        with self.assertRaises(ValueError):
            feishu_sheets.sync(self.folder, self.config)
        path = self.folder / "feishu_sheet.json"
        saved = json.loads(path.read_text())
        saved["create_title"] += " · legacy-id"
        path.write_text(json.dumps(saved))
        self.server.sheets[-1]["title"] = saved["create_title"]
        self.assertEqual(feishu_sheets.sync(self.folder, self.config)["status"], "complete")
        self.assertEqual(len(self.creates()), 1)
        self.assertEqual(self.server.sheets[-1]["title"], saved["create_title"])

    def test_partial_write_retries_from_unconfirmed_range(self):
        end = feishu_sheets.get_column_letter(1 + len(feishu_sheets.COLUMNS))
        self.server.fail_range = f"tab1!A3:{end}4"
        with self.assertRaisesRegex(ValueError, "写入请求超时"):
            feishu_sheets.sync(self.folder, self.config)
        self.assertEqual(feishu_sheets.saved_state(self.folder)["written_rows"], 2)
        first_batch_calls = len([1 for _, path, data in self.server.calls
                                 if path.endswith("/values") and data["valueRange"]["range"] == f"tab1!A1:{end}2"])
        feishu_sheets.sync(self.folder, self.config)
        self.assertEqual(len(self.creates()), 1)
        self.assertEqual(len([1 for _, path, data in self.server.calls
                              if path.endswith("/values") and data["valueRange"]["range"] == f"tab1!A1:{end}2"]), first_batch_calls)

    def restore_previous_export_layout(self, with_suggestion):
        from openpyxl import load_workbook
        from openpyxl.worksheet.table import TableColumn
        book = load_workbook(self.folder / "feedback_details.xlsx")
        for sheet in book:
            sheet.insert_cols(9, 2)
            sheet.cell(4, 9, "原标签")
            sheet.cell(4, 10, "用户主动选择的标签")
            for row in range(5, sheet.max_row + 1):
                sheet.cell(row, 9, "旧原标签")
                sheet.cell(row, 10, "旧玩家标签")
            if not with_suggestion:
                sheet.delete_cols(12)
            for table in sheet.tables.values():
                table.ref = f"A4:{feishu_sheets.get_column_letter(sheet.max_column)}{sheet.max_row}"
                table.autoFilter.ref = table.ref
                table.tableColumns = [TableColumn(id=index, name=cell.value)
                                      for index, cell in enumerate(sheet[4], 1)]
        book.save(self.folder / "feedback_details.xlsx")
        book.close()

    def test_legacy_report_keeps_original_columns_and_can_resume(self):
        self.restore_previous_export_layout(with_suggestion=False)
        rows = feishu_sheets.detail_rows(self.folder)
        self.assertEqual(len(rows[0]), 12)
        self.assertNotIn("是否建议", rows[0])
        self.assertEqual(rows[1][3], "004")
        self.assertEqual(rows[1][8], "'=1+1")
        self.assertEqual(feishu_sheets.sync(self.folder, self.config)["status"], "complete")
        count = len(self.server.calls)
        self.assertEqual(feishu_sheets.sync(self.folder, self.config)["status"], "complete")
        self.assertEqual(len(self.server.calls), count)

    def test_previous_suggestion_export_keeps_removed_columns_and_can_resume(self):
        self.restore_previous_export_layout(with_suggestion=True)
        rows = feishu_sheets.detail_rows(self.folder)
        self.assertEqual(rows[0], ["统计期间", *feishu_sheets.PREVIOUS_DETAIL_COLUMNS])
        self.assertEqual(rows[1][9:11], ["旧原标签", "旧玩家标签"])
        self.assertEqual(rows[1][-1], "是")
        self.assertEqual(feishu_sheets.sync(self.folder, self.config)["status"], "complete")
        count = len(self.server.calls)
        self.assertEqual(feishu_sheets.sync(self.folder, self.config)["status"], "complete")
        self.assertEqual(len(self.server.calls), count)

    def test_date_format_retains_time_and_uses_verified_api_format(self):
        feishu_sheets.sync(self.folder, self.config)
        formats = [item for _, path, payload in self.server.calls if path.endswith("/styles_batch_update")
                   for item in payload["data"] if "formatter" in item["style"]]
        dates = [item for item in formats if item["ranges"][0].startswith("tab1!G")]
        self.assertTrue(dates)
        self.assertTrue(all(item["style"]["formatter"] == "yyyy/MM/dd HH:mm:ss" for item in dates))
        self.assertIn("tab1!G2:G2", dates[0]["ranges"])

    def test_retry_after_invalid_format_reuses_existing_empty_tab(self):
        original = self.server.call
        failed = False
        def reject_first_batch(method, path, payload=None):
            nonlocal failed
            if path.endswith("/styles_batch_update") and not failed and any(
                    "formatter" in item["style"] for item in payload["data"]):
                failed = True
                raise ApiError("invalid formatter", 90204)
            return original(method, path, payload)
        with patch.object(self.server, "call", side_effect=reject_first_batch):
            with self.assertRaisesRegex(ValueError, "invalid formatter"):
                feishu_sheets.sync(self.folder, self.config)
        saved = feishu_sheets.saved_state(self.folder)
        self.assertTrue(saved["prepared"])
        self.assertEqual(saved["written_rows"], 0)
        self.assertEqual(feishu_sheets.sync(self.folder, self.config)["status"], "complete")
        self.assertEqual(len(self.creates()), 1)

    def test_changed_target_is_not_silently_overwritten(self):
        feishu_sheets.sync(self.folder, self.config)
        self.config["feishu"]["spreadsheet"]["url"] = "https://example.feishu.cn/sheets/another"
        with self.assertRaisesRegex(ValueError, "已经变化"):
            feishu_sheets.sync(self.folder, self.config)

    def test_explicit_token_skips_wiki_permissions(self):
        self.config["feishu"]["spreadsheet"]["spreadsheet_token"] = "target"
        feishu_sheets.sync(self.folder, self.config)
        self.assertFalse(any("/wiki/" in path for _, path, _ in self.server.calls))

    def test_card_waits_for_details_then_includes_tab_link(self):
        with patch("feedback.feishu.credentials", return_value=("https://example", "")), patch(
                "feedback.feishu.send_payload", return_value={"code": 0}) as sender:
            result = feishu.push(self.folder, self.config)
            text = sender.call_args.args[1]["card"]["body"]["elements"][0]["content"]
            self.assertIn("打开完整反馈明细", text)
            self.assertIn("sheet=tab1", text)
            self.assertEqual(result["sheet"]["status"], "complete")
            feishu.push(self.folder, self.config)
            sender.assert_called_once()

    def test_details_failure_does_not_send_card(self):
        self.server.fail_create = True
        with patch("feedback.feishu.credentials", return_value=("https://example", "")), patch(
                "feedback.feishu.send_payload") as sender:
            with self.assertRaises(ValueError):
                feishu.push(self.folder, self.config)
            sender.assert_not_called()

    def test_large_export_enlarges_only_new_tab(self):
        rows = [["统计期间", *feishu_sheets.COLUMNS.values()]]
        rows += [["本期", *(["示例"] * len(feishu_sheets.COLUMNS))] for _ in range(505)]
        with patch("feedback.feishu_sheets.detail_rows", return_value=rows):
            feishu_sheets.sync(self.folder, self.config)
        dimensions = [data["dimension"] for _, path, data in self.server.calls if path.endswith("/dimension_range")]
        self.assertEqual(dimensions, [{"sheetId": "tab1", "majorDimension": "ROWS", "length": 306}])
        self.assertEqual(self.server.sheets[0]["grid_properties"]["row_count"], 200)
        self.assertTrue(all(len(batch) <= 2 for batch in self.server.writes.values()))

    def test_incomplete_write_response_is_not_marked_complete(self):
        original = self.server.call
        def response(method, path, payload=None):
            data = original(method, path, payload)
            if path.endswith("/values"):
                data["updatedRows"] = 0
            return data
        with patch.object(self.server, "call", side_effect=response):
            with self.assertRaisesRegex(ValueError, "行列数不匹配"):
                feishu_sheets.sync(self.folder, self.config)
        self.assertEqual(feishu_sheets.saved_state(self.folder)["written_rows"], 0)

    def test_wrong_wiki_resource_type_is_rejected_before_creation(self):
        with patch.object(self.server, "call", return_value={"node": {"obj_type": "docx", "obj_token": "not-a-sheet"}}):
            with self.assertRaisesRegex(ValueError, "不是飞书电子表格"):
                feishu_sheets.sync(self.folder, self.config)
        self.assertEqual(len(self.server.sheets), 1)


class AuthenticationTests(BaseTest):
    def test_token_cache_uses_expiry_and_never_saves_secret(self):
        with patch.dict(os.environ, {"TEST_ID": "app", "TEST_SECRET": "secret"}), patch(
                "feedback.feishu_api.request_json", return_value={"code": 0, "tenant_access_token": "token", "expire": 7200}) as request:
            client = Client({"app_id_env": "TEST_ID", "app_secret_env": "TEST_SECRET"})
            _tokens.pop(client.cache_key, None)
            self.assertEqual(client.token(), "token")
            self.assertEqual(client.token(), "token")
            request.assert_called_once()
            self.assertNotIn("secret", client.cache_key)

    def test_invalid_or_non_feishu_destination_is_rejected(self):
        for url in ("http://example.feishu.cn/sheets/test", "https://evil.com/wiki/test", "https://example.feishu.cn/base/test"):
            with self.assertRaises(ValueError):
                feishu_sheets.destination({"url": url})
