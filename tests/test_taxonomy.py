from copy import deepcopy
from unittest.mock import patch

from openpyxl import load_workbook

from feedback.analysis import analyze_report
from feedback.cleaning import clean_feedback
from feedback.config import load_config, read_json
from feedback.feishu_sheets import detail_rows
from feedback.processing import process_feedback, validate_result
from feedback.selection import select_feedback
from feedback.statistics import build_statistics
from feedback.storage import atomic_csv, read_csv
from feedback.taxonomy import display_theme
from feedback.workflow import run_report
from tests.helpers import BaseTest, fake_request


def answer(**changes):
    return {"translation_zh": "玩家反馈内容。", "label": "玩法相关", "theme": "活动", "subtheme": "许愿树",
            "is_valid": True, "exclusion_reason": "", **changes}


class ThreeLevelTaxonomyTests(BaseTest):
    def selected(self):
        cleaned, _ = clean_feedback(self.raw, self.config)
        return select_feedback(self.raw, cleaned, "2026-10-19", "2026-10-22", self.config)

    def processed(self):
        rows, info = self.selected()
        rows["最终标签"], rows["具体主题"], rows["三级主题"] = "玩法相关", "活动", "许愿树"
        rows["中文翻译"], rows["是否有效"], rows["排除原因"] = "玩家咨询活动。", "是", ""
        return rows, info

    def test_six_roots_and_confirmed_relocations(self):
        labels = self.config["schema"]["labels"]
        self.assertEqual(list(labels), ["账号问题", "游戏卡顿", "玩法相关", "充值", "作弊", "其他"])
        gameplay = {node["name"]: node for node in labels["玩法相关"]["themes"]}
        self.assertEqual({node["name"] for node in gameplay["活动"]["children"]},
                         {"无尽礼包", "存钱袋", "幸运翻倍", "周末礼包", "许愿树", "沙漠寻宝"})
        self.assertIn("其他玩法限制", {node["name"] for node in gameplay["游戏限制"]["children"]})
        self.assertFalse(gameplay["俱乐部"].get("children"))
        self.assertIn("其他充值问题", {node["name"] for node in labels["充值"]["themes"]})
        other = {node["name"]: node for node in labels["其他"]["themes"]}
        self.assertTrue({"封禁", "个人信息页面", "相册功能", "金币丢失", "抱怨"} <= set(other))
        self.assertFalse(other["个人信息页面"].get("children"))

    def test_valid_paths_keep_three_separate_fields(self):
        for label, theme, subtheme in [("玩法相关", "活动", "许愿树"), ("玩法相关", "游戏限制", "其他玩法限制"),
                ("玩法相关", "发牌&公平", "发牌不公平"), ("玩法相关", "段位与赛季", "赛季结算&赛季奖励"),
                ("玩法相关", "俱乐部", ""), ("玩法相关", "翻倍场", ""), ("玩法相关", "活动", ""),
                ("其他", "封禁", ""), ("其他", "个人信息页面", ""), ("其他", "相册功能", ""),
                ("充值", "其他充值问题", ""), ("游戏卡顿", "", "")]:
            with self.subTest(path=(label, theme, subtheme)):
                result = validate_result(answer(label=label, theme=theme, subtheme=subtheme), self.config)
                self.assertEqual((result["label"], result["theme"], result["subtheme"]), (label, theme, subtheme))
        self.assertEqual(display_theme("活动", "许愿树"), "活动（许愿树）")
        self.assertEqual(display_theme("俱乐部", ""), "俱乐部")

    def test_wrong_parent_flattened_or_legacy_paths_are_rejected(self):
        for changes in [{"label": "活动", "theme": "许愿树", "subtheme": ""},
                {"subtheme": "欢乐场游戏限制"}, {"label": "充值"},
                {"theme": "翻倍场"}, {"theme": ""}, {"theme": "活动（许愿树）", "subtheme": ""},
                {"label": "其他", "theme": "个人信息页面", "subtheme": "战绩记录"}]:
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                validate_result(answer(**changes), self.config)
        legacy = answer()
        legacy.pop("subtheme")
        with self.assertRaisesRegex(ValueError, "三层分类"):
            validate_result(legacy, self.config)
        with self.assertRaises(ValueError):
            validate_result(answer(label="", theme="", is_valid=False, exclusion_reason="含辱骂内容"), self.config)

    def test_malformed_or_deeper_taxonomies_are_rejected_at_load(self):
        bad_children = [[{"name": "重复"}, {"name": "重复"}], [{"name": "第四层", "children": []}], "invalid"]
        for children in bad_children:
            def read(path):
                value = read_json(path)
                if str(path).endswith("baloot/labels.json"):
                    value["labels"]["玩法相关"]["themes"].append({"name": "测试分组", "children": deepcopy(children)})
                return value
            with self.subTest(children=children), patch("feedback.config.read_json", side_effect=read):
                with self.assertRaises(ValueError):
                    load_config("baloot")

    def test_workflow_exports_two_display_columns_and_preserves_hierarchy_in_cache(self):
        paths = {"001": ("玩法相关", "活动", "沙漠寻宝"), "002": ("充值", "其他充值问题", ""),
                 "004": ("玩法相关", "活动", "许愿树"), "005": ("玩法相关", "俱乐部", "")}

        def request(prompt, payload, config, validate, should_stop=None):
            if "feedbacks" not in payload:
                return fake_request(prompt, payload, config, validate, should_stop)
            return validate({"results": [dict(feedback_id=item["feedback_id"], **answer(
                label=paths[item["feedback_id"]][0], theme=paths[item["feedback_id"]][1],
                subtheme=paths[item["feedback_id"]][2])) for item in payload["feedbacks"]]})

        with patch("feedback.ai_client.request_json", side_effect=request) as model:
            result = run_report(self.root, self.config, "2026-10-19", "2026-10-22", False, "hierarchy", lambda **v: None)
            calls = model.call_count
            again = run_report(self.root, self.config, "2026-10-19", "2026-10-22", False, "cached", lambda **v: None)
            self.assertEqual(model.call_count, calls)
            self.assertEqual(sum("feedbacks" in call.args[1] for call in model.call_args_list), 1)
        self.assertEqual((result["total"], result["previous_total"]), (2, 2))
        self.assertEqual(again, result)
        stored = read_csv(self.root / "processed.csv").set_index("反馈ID")
        self.assertEqual(stored.loc["004", "三级主题"], "许愿树")
        self.assertEqual(stored.loc["005", "三级主题"], "")
        folder = self.root / "reports/hierarchy"
        book = load_workbook(folder / "feedback_details.xlsx", read_only=True)
        try:
            for sheet in book:
                self.assertEqual([cell.value for cell in sheet[4]][:2], ["I 类别", "II 类别"])
                self.assertEqual(sheet.max_column, 9)
                self.assertNotIn("是否建议", [cell.value for cell in sheet[4]])
            values = {row[2]: row[:2] for row in book["本期反馈明细"].iter_rows(min_row=5, values_only=True)}
            self.assertEqual(values["004"], ("玩法相关", "活动（许愿树）"))
            self.assertEqual(values["005"], ("玩法相关", "俱乐部"))
        finally:
            book.close()
        feishu = detail_rows(folder)
        self.assertEqual(feishu[0][:3], ["统计期间", "I 类别", "II 类别"])
        self.assertEqual(next(row[2] for row in feishu[1:] if row[3] == "004"), "活动（许愿树）")

    def test_issue_analysis_keeps_third_level_counts_separate_under_one_root(self):
        rows, info = self.processed()
        rows.loc[rows["反馈ID"].eq("005"), "三级主题"] = "沙漠寻宝"
        self.config["report"]["module_growth_thresholds"]["玩法相关"] = 0
        snapshot = build_statistics(rows, info, self.config)
        self.assertEqual(len(snapshot["modules"]), 1)
        self.assertEqual(snapshot["modules"][0]["count"], 2)
        payloads = []

        def request(name, payload, config, cache, validate, should_stop=None):
            self.assertEqual(name, "issues")
            payloads.append(payload)
            return validate({"text": "玩家咨询该活动的规则。", "feedback_ids": [payload["items"][0]["feedback_id"]]})

        with patch("feedback.analysis.cached_request", side_effect=request):
            result = analyze_report(snapshot, self.config, self.root / "cache", lambda **v: None)
        self.assertEqual({issue["theme"]: issue["count"] for issue in result["issues"]["玩法相关"]},
                         {"活动（许愿树）": 1, "活动（沙漠寻宝）": 1})
        self.assertTrue(all(len({item["subtheme"] for item in payload["items"]}) == 1 for payload in payloads))

    def test_suggestion_consolidation_retains_second_and_third_level_context(self):
        rows, info = self.processed()
        rows["反馈标签"] = "建议"
        rows.loc[rows["反馈ID"].eq("005"), ["具体主题", "三级主题"]] = ["俱乐部", ""]
        snapshot = build_statistics(rows, info, self.config)

        def request(prompt, payload, config, validate, should_stop=None):
            if payload["stage"] == "提取":
                return validate({"suggestions": [{"text": "增加入口", "feedback_ids": [item["feedback_id"]]}
                    for item in payload["items"]]})
            self.assertEqual({tuple(item["categories"]) for item in payload["items"]},
                             {("活动（许愿树）",), ("俱乐部",)})
            return validate({"suggestions": [{"text": "增加入口", "candidate_ids": [item["candidate_id"]]}
                for item in payload["items"]]})

        with patch("feedback.ai_client.request_json", side_effect=request):
            result = analyze_report(snapshot, self.config, self.root / "cache", lambda **v: None)
        self.assertEqual(len(result["suggestions"]), 2)

    def test_legacy_cache_missing_third_level_is_reprocessed(self):
        rows, _ = self.selected()
        path = self.root / "processed.csv"
        with patch("feedback.ai_client.request_json", side_effect=fake_request):
            process_feedback(rows, self.config, path, lambda **v: None)
        atomic_csv(path, read_csv(path).drop(columns=["三级主题"]))
        with patch("feedback.ai_client.request_json", side_effect=fake_request) as model:
            result = process_feedback(rows, self.config, path, lambda **v: None)
        self.assertEqual(model.call_count, 1)
        self.assertTrue(result["三级主题"].eq("").all())
