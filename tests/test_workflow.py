from copy import deepcopy
from unittest.mock import patch
from openpyxl import load_workbook
from feedback.workflow import run_report
from feedback.storage import read_csv
from tests.helpers import BaseTest, fake_request


class WorkflowTests(BaseTest):
    def test_changing_growth_threshold_reuses_classification_and_only_adds_required_analysis(self):
        first = deepcopy(self.config)
        first["report"]["module_growth_thresholds"]["玩法相关"] = 2

        def request(prompt, payload, config, validate, should_stop=None):
            if "feedbacks" not in payload:
                return fake_request(prompt, payload, config, validate, should_stop)
            return validate({"results": [{"feedback_id": item["feedback_id"], "translation_zh": "玩家咨询玩法。",
                "label": "游戏卡顿" if item["feedback_id"] == "002" else "玩法相关",
                "theme": "" if item["feedback_id"] == "002" else "翻倍场", "subtheme": "",
                "is_valid": True, "exclusion_reason": ""} for item in payload["feedbacks"]]})

        with patch("feedback.ai_client.request_json", side_effect=request) as model:
            before = run_report(self.root, first, "2026-10-19", "2026-10-22", False, "high-threshold", lambda **v: None)
            self.assertEqual(model.call_count, 1)
            lowered = deepcopy(first)
            lowered["report"]["module_growth_thresholds"]["玩法相关"] = .1
            after = run_report(self.root, lowered, "2026-10-19", "2026-10-22", False, "low-threshold", lambda **v: None)
            self.assertEqual(sum("feedbacks" in call.args[1] for call in model.call_args_list), 1)
            self.assertEqual(sum("items" in call.args[1] for call in model.call_args_list), 1)
            count = model.call_count
            run_report(self.root, lowered, "2026-10-19", "2026-10-22", False, "same-threshold", lambda **v: None)
            self.assertEqual(model.call_count, count)
        self.assertEqual((before["total"], before["previous_total"]), (after["total"], after["previous_total"]))

    def test_whole_workflow_selects_four_rows_and_exports_two_sheets(self):
        with patch("feedback.ai_client.request_json",side_effect=fake_request) as request:
            result = run_report(self.root,self.config,"2026-10-19","2026-10-22",False,"first",lambda **v:None)
            translation_calls = [c for c in request.call_args_list if "feedbacks" in c.args[1]]
            self.assertEqual(len(translation_calls),1)
            self.assertEqual(len(translation_calls[0].args[1]["feedbacks"]),4)
            self.assertEqual(result["total"],2)
            self.assertEqual(result["previous_total"],2)
            calls = request.call_count
            run_report(self.root,self.config,"2026-10-19","2026-10-22",False,"second",lambda **v:None)
            self.assertEqual(request.call_count,calls)
        folder = self.root / "reports/first"
        book = load_workbook(folder/"feedback_details.xlsx")
        self.assertEqual(book.sheetnames,["本期反馈明细","上周同期反馈明细"])
        self.assertEqual(book.worksheets[0].max_row,6)
        cell = book.worksheets[0]["H5"]
        self.assertEqual(cell.value,"=1+1")
        self.assertEqual(cell.data_type,"s")
        self.assertEqual(book.worksheets[0]["C5"].value,"004")
        book.close()
        self.assertIn("2026-10-12 至 2026-10-15",(folder/"report.md").read_text())
        self.assertEqual(len(read_csv(self.root/"processed.csv")),4)

    def test_failure_does_not_publish_report(self):
        with patch("feedback.ai_client.request_json",side_effect=ValueError("test failure")):
            with self.assertRaises(ValueError):
                run_report(self.root,self.config,"2026-10-19","2026-10-22",False,"failed",lambda **v:None)
        self.assertFalse((self.root/"reports/failed/report.md").exists())

    def test_analysis_cannot_invent_evidence_ids(self):
        from feedback.analysis import evidence
        with self.assertRaises(ValueError):
            evidence({"text":"summary","feedback_ids":["invented"]},{"001"},self.config,160)
