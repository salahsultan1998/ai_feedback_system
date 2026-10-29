from unittest.mock import patch

from openpyxl import load_workbook

from feedback.analysis import analyze_report
from feedback.cleaning import clean_feedback
from feedback.feishu_sheets import detail_rows
from feedback.processing import fingerprint, process_feedback, validate_result
from feedback.policies import ABUSE_REASON
from feedback.selection import select_feedback
from feedback.statistics import build_statistics
from feedback.storage import atomic_csv, read_csv
from feedback.workflow import run_report
from tests.helpers import BaseTest, fake_request


def answer(**changes):
    return {"translation_zh": "玩家反馈内容。", "label": "玩法相关", "theme": "", "subtheme": "",
            "is_valid": True, "exclusion_reason": "", **changes}


class ClassificationQualityTests(BaseTest):
    def selected(self, start="2026-10-19", end="2026-10-22"):
        cleaned, _ = clean_feedback(self.raw, self.config)
        return select_feedback(self.raw, cleaned, start, end, self.config)

    def processed(self, start="2026-10-19", end="2026-10-22"):
        selected, info = self.selected(start, end)
        selected["最终标签"], selected["具体主题"], selected["三级主题"], selected["中文翻译"] = "玩法相关", "俱乐部", "", "玩家反馈俱乐部问题。"
        selected["是否有效"], selected["是否建议"], selected["排除原因"] = "是", "否", ""
        return selected, info

    def test_suggestion_comes_from_source_label_and_cannot_be_a_module_or_theme(self):
        for label, theme in [("玩法相关", "俱乐部"), ("玩法相关", "欢乐场"), ("其他", ""),
                             ("其他", "相册功能"), ("充值", ""), ("玩法相关", "活动")]:
            for suggested in (True, False):
                with self.subTest(label=label, suggested=suggested):
                    result = validate_result(answer(label=label, theme=theme, is_suggestion=not suggested),
                                             self.config, "建议" if suggested else "投诉")
                    self.assertEqual(result["label"], label)
                    self.assertIs(result["is_suggestion"], suggested)
        for value in (answer(label="建议"), answer(label="其他", theme="建议")):
            with self.assertRaises(ValueError):
                validate_result(value, self.config)

    def test_flags_and_exclusion_decisions_are_strict(self):
        for field in ("is_valid",):
            for bad in ("false", "true", 0, 1, None):
                with self.subTest(field=field, bad=bad), self.assertRaises(ValueError):
                    validate_result(answer(**{field: bad}), self.config)
        invalid = answer(is_valid=False, label="", exclusion_reason="缺少明确反馈对象或诉求")
        self.assertFalse(validate_result(invalid, self.config)["is_valid"])
        for changes in ({"label": "其他"}, {"exclusion_reason": ""},
                        {"exclusion_reason": "模型自创原因"}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                validate_result({**invalid, **changes}, self.config)
        with self.assertRaises(ValueError):
            validate_result(answer(exclusion_reason="内容无法理解"), self.config)
        self.assertFalse(validate_result({**invalid, "is_suggestion": True}, self.config, "建议")["is_suggestion"])

    def test_source_label_matching_is_exact_and_model_flags_cannot_override_it(self):
        for source, expected in [("建议", True), (" 建议 ", True), ("非建议", False), ("建议性投诉", False),
                                 ("投诉", False), ("", False)]:
            for returned in (True, False, "true", "false", None):
                with self.subTest(source=source, returned=returned):
                    self.assertIs(validate_result(answer(is_suggestion=returned), self.config, source)["is_suggestion"], expected)
        configured = {**self.config, "suggestion_labels": ["产品建议"]}
        self.assertTrue(validate_result(answer(), configured, "产品建议")["is_suggestion"])
        self.assertFalse(validate_result(answer(), configured, "建议")["is_suggestion"])

    def test_module_counts_do_not_double_count_suggestions_and_exclude_ambiguous_rows(self):
        rows, info = self.processed()
        rows.loc[rows["反馈ID"].eq("004"), "反馈标签"] = "建议"
        rows.loc[rows["反馈ID"].eq("005"), "反馈标签"] = "建议"
        rows.loc[rows["反馈ID"].eq("005"), ["是否有效", "最终标签", "排除原因"]] = ["否", "", "缺少明确反馈对象或诉求"]
        result = build_statistics(rows, info, self.config)
        self.assertEqual((result["total"], result["previous_total"]), (1, 2))
        self.assertEqual(result["suggestion_feedback_count"], 1)
        self.assertEqual(sum(m["count"] for m in result["modules"]), 1)
        self.assertEqual(sum(m["share"] for m in result["modules"]), 1)
        self.assertEqual(result["modules"][0]["suggestion_count"], 1)
        self.assertEqual(result["selection"]["candidate_current_count"], 2)
        self.assertEqual(result["selection"]["semantic_excluded_current"], 1)

    def test_suggestion_extraction_uses_only_source_tagged_records_across_modules(self):
        rows, info = self.processed()
        rows.loc[rows["反馈ID"].eq("004"), ["最终标签", "具体主题", "是否建议"]] = ["其他", "", "是"]
        rows.loc[rows["反馈ID"].eq("004"), "反馈标签"] = "建议"
        rows.loc[rows["反馈ID"].eq("005"), ["最终标签", "具体主题"]] = ["其他", "个人信息页面"]
        snapshot = build_statistics(rows, info, self.config)
        # A stale/model-supplied flag must not move a record into or out of the pool.
        snapshot["current"].loc[:, "是否建议"] = ["否", "是"]
        for module in snapshot["modules"]:
            module["flagged"] = False  # Isolate the suggestion-pool check from issue analysis.
        payloads = []
        def request(name, payload, config, cache, validate, should_stop=None):
            self.assertEqual(name, "suggestions")
            payloads.append(payload)
            return validate({"suggestions": [{"text": "建议增加一个新玩法。", "feedback_ids": ["004"]}]})
        with patch("feedback.analysis.cached_request", side_effect=request):
            result = analyze_report(snapshot, self.config, self.root / "cache", lambda **v: None)
        self.assertEqual([i["feedback_id"] for p in payloads for i in p["items"]], ["004"])
        self.assertTrue(all(i["is_suggestion"] for p in payloads for i in p["items"]))
        self.assertEqual(result["suggestions"][0]["feedback_ids"], ["004"])

    def test_filter_result_is_cached_and_exported_without_entering_counts_or_feishu(self):
        submissions = []
        def request(prompt, payload, config, validate, should_stop=None):
            if "feedbacks" not in payload:
                return fake_request(prompt, payload, config, validate, should_stop)
            submissions.extend(i["feedback_id"] for i in payload["feedbacks"])
            return validate({"results": [dict(feedback_id=i["feedback_id"], **(
                answer(label="", is_valid=False, exclusion_reason="缺少明确反馈对象或诉求")
                if i["feedback_id"] == "004" else answer())) for i in payload["feedbacks"]]})
        with patch("feedback.ai_client.request_json", side_effect=request):
            result = run_report(self.root, self.config, "2026-10-19", "2026-10-22", False, "filtered", lambda **v: None)
            run_report(self.root, self.config, "2026-10-19", "2026-10-22", False, "reused", lambda **v: None)
        self.assertEqual(sorted(submissions), ["001", "002", "004", "005"])
        self.assertEqual((result["total"], result["previous_total"]), (1, 2))
        stored = read_csv(self.root / "processed.csv").set_index("反馈ID")
        self.assertEqual(stored.loc["004", "是否有效"], "否")
        self.assertEqual(stored.loc["004", "错误"], "")
        folder = self.root / "reports/filtered"
        book = load_workbook(folder / "feedback_details.xlsx")
        try:
            self.assertEqual(book.sheetnames, ["本期反馈明细", "上周同期反馈明细", "排除反馈明细"])
            self.assertEqual(book["本期反馈明细"].max_column, 9)
            for sheet in book:
                headers = [cell.value for cell in sheet[4]]
                self.assertNotIn("原标签", headers)
                self.assertNotIn("用户主动选择的标签", headers)
                self.assertNotIn("是否建议", headers)
            self.assertEqual(book["本期反馈明细"].column_dimensions["I"].width, 40)
            self.assertEqual(stored.loc["005", "是否建议"], "否")
            self.assertEqual(book["排除反馈明细"]["A5"].value, "004")
            self.assertEqual(book["排除反馈明细"]["E5"].value, "=1+1")
            self.assertEqual(book["排除反馈明细"]["E5"].data_type, "s")
            self.assertEqual(book["排除反馈明细"]["F5"].value, "缺少明确反馈对象或诉求")
        finally:
            book.close()
        self.assertNotIn("004", {r[3] for r in detail_rows(folder)[1:]})
        markdown = (folder / "report.md").read_text()
        self.assertIn("本期有效反馈：1 条", markdown)
        self.assertEqual(result["selection"]["semantic_excluded_current"], 1)
        self.assertEqual(result["selection"]["semantic_excluded_previous"], 0)

    def test_zero_current_after_filtering_and_overlapping_periods_are_supported(self):
        rows, info = self.processed()
        mask = rows["反馈ID"].isin(["004", "005"])
        rows.loc[mask, ["是否有效", "最终标签", "排除原因"]] = ["否", "", "内容无法理解"]
        result = build_statistics(rows, info, self.config)
        self.assertEqual((result["total"], result["previous_total"]), (0, 2))
        self.assertEqual(result["modules"][0]["share"], 0)
        with patch("feedback.analysis.cached_request") as request:
            self.assertEqual(analyze_report(result, self.config, self.root / "cache", lambda **v: None),
                             {"suggestions": [], "issues": {}})
            request.assert_not_called()
        rows, info = self.processed("2026-10-12", "2026-10-22")
        rows.loc[rows["反馈ID"].eq("002"), ["是否有效", "最终标签", "排除原因"]] = ["否", "", "内容无法理解"]
        result = build_statistics(rows, info, self.config)
        self.assertEqual(result["selection"]["semantic_excluded_in_period"], 1)
        self.assertEqual(result["selection"]["semantic_excluded_current"], 1)
        self.assertEqual(result["selection"]["semantic_excluded_previous"], 1)

    def test_old_cache_cannot_guess_missing_suggestion_flags(self):
        rows, _ = self.selected()
        rows["反馈标签"] = "建议"
        cached = rows.copy()
        cached["中文翻译"], cached["最终标签"], cached["具体主题"], cached["错误"] = "旧翻译", "玩法相关", "俱乐部", ""
        cached["处理指纹"] = [fingerprint(row, self.config) for _, row in rows.iterrows()]
        path = self.root / "processed.csv"
        atomic_csv(path, cached)
        with patch("feedback.ai_client.request_json", side_effect=fake_request) as request:
            result = process_feedback(rows, self.config, path, lambda **v: None)
        self.assertEqual(request.call_count, 1)
        self.assertTrue(result["是否建议"].eq("是").all())

    def test_all_current_excluded_can_publish_a_zero_count_report(self):
        def request(prompt, payload, config, validate, should_stop=None):
            self.assertIn("feedbacks", payload, "No report analysis is needed with zero valid current records")
            return validate({"results": [dict(feedback_id=i["feedback_id"], **(
                answer(label="", is_valid=False, exclusion_reason="内容无法理解")
                if i["feedback_id"] in {"004", "005"} else answer())) for i in payload["feedbacks"]]})
        with patch("feedback.ai_client.request_json", side_effect=request):
            result = run_report(self.root, self.config, "2026-10-19", "2026-10-22", False, "empty-current", lambda **v: None)
        self.assertEqual((result["total"], result["previous_total"], result["suggestions"]), (0, 2, 0))
        folder = self.root / "reports/empty-current"
        book = load_workbook(folder / "feedback_details.xlsx")
        try:
            self.assertEqual(book["本期反馈明细"].max_row, 4)
            self.assertEqual(book["上周同期反馈明细"].max_row, 6)
            self.assertEqual(book["排除反馈明细"].max_row, 6)
        finally:
            book.close()
        self.assertEqual([row[0] for row in detail_rows(folder)[1:]], ["上周同期", "上周同期"])

    def test_abusive_suggestions_and_complaints_are_removed_from_every_report_output(self):
        self.raw.loc[self.raw["反馈ID"].eq("004"), "用户自己填写的内容"] = "You idiots, please add a club ranking."
        self.raw.loc[self.raw["反馈ID"].eq("001"), "用户自己填写的内容"] = "Stupid game, I cannot recharge."
        self.raw.loc[self.raw["反馈ID"].isin(["004", "005"]), "反馈标签"] = "建议"
        atomic_csv(self.root / "rows.csv", self.raw)
        analysis_ids = []

        def request(prompt, payload, config, validate, should_stop=None):
            if "feedbacks" in payload:
                results = []
                for item in payload["feedbacks"]:
                    fid = item["feedback_id"]
                    result = (answer(label="", is_valid=False, exclusion_reason=ABUSE_REASON)
                              if fid in {"001", "004"} else
                              answer(label="", is_valid=False, exclusion_reason="缺少明确反馈对象或诉求")
                              if fid == "002" else answer())
                    results.append(dict(feedback_id=fid, **result))
                return validate({"results": results})
            analysis_ids.extend(item["feedback_id"] for item in payload["items"])
            return fake_request(prompt, payload, config, validate, should_stop)

        with patch("feedback.ai_client.request_json", side_effect=request) as model:
            result = run_report(self.root, self.config, "2026-10-19", "2026-10-22", False, "no-abuse", lambda **v: None)
            count = model.call_count
            run_report(self.root, self.config, "2026-10-19", "2026-10-22", False, "resumed", lambda **v: None)
            self.assertEqual(model.call_count, count)
        self.assertEqual((result["total"], result["previous_total"], result["suggestion_feedback_count"]), (1, 0, 1))
        self.assertEqual(result["selection"]["abusive_excluded_in_period"], 2)
        self.assertEqual(result["selection"]["abusive_excluded_current"], 1)
        self.assertEqual(result["selection"]["abusive_excluded_previous"], 1)
        self.assertEqual(result["selection"]["semantic_excluded_in_period"], 1)
        self.assertEqual(set(analysis_ids), {"005"})
        stored = read_csv(self.root / "processed.csv").set_index("反馈ID")
        self.assertTrue(stored.loc[["001", "004"], "是否建议"].eq("否").all())
        self.assertTrue(stored.loc[["001", "004"], "错误"].eq("").all())
        folder = self.root / "reports/no-abuse"
        book = load_workbook(folder / "feedback_details.xlsx")
        try:
            self.assertEqual(book["排除反馈明细"]["A5"].value, "002")
            for sheet in book:
                headers = [cell.value for cell in sheet[4]]
                ids = [row[headers.index("反馈ID")] for row in sheet.iter_rows(min_row=5, values_only=True)]
                self.assertFalse({"001", "004"} & set(ids))
        finally:
            book.close()
        self.assertEqual([r[3] for r in detail_rows(folder)[1:]], ["005"])
        markdown = (folder / "report.md").read_text()
        self.assertNotIn("001", markdown)
        self.assertNotIn("004", markdown)
        self.assertIn("（用户 ID：0009）", markdown)

    def test_all_abusive_feedback_exports_empty_valid_sheets_without_excluded_sheet(self):
        def request(prompt, payload, config, validate, should_stop=None):
            self.assertIn("feedbacks", payload)
            return validate({"results": [dict(feedback_id=item["feedback_id"], **answer(
                label="", is_valid=False, exclusion_reason=ABUSE_REASON)) for item in payload["feedbacks"]]})
        with patch("feedback.ai_client.request_json", side_effect=request):
            result = run_report(self.root, self.config, "2026-10-19", "2026-10-22", False, "all-abusive", lambda **v: None)
        self.assertEqual((result["total"], result["previous_total"], result["suggestions"]), (0, 0, 0))
        self.assertEqual(result["selection"]["abusive_excluded_in_period"], 4)
        self.assertEqual(result["selection"]["semantic_excluded_in_period"], 0)
        folder = self.root / "reports/all-abusive"
        book = load_workbook(folder / "feedback_details.xlsx")
        try:
            self.assertEqual(book.sheetnames, ["本期反馈明细", "上周同期反馈明细"])
            self.assertTrue(all(sheet.max_row == 4 for sheet in book))
        finally:
            book.close()
        self.assertEqual(len(detail_rows(folder)), 1)

    def test_abuse_in_overlapping_periods_is_excluded_once_from_union(self):
        rows, info = self.processed("2026-10-12", "2026-10-22")
        # The fallback also rejects a stale valid row classified into the retired module.
        rows.loc[rows["反馈ID"].eq("002"), "最终标签"] = "辱骂"
        result = build_statistics(rows, info, self.config)
        self.assertEqual(result["selection"]["abusive_excluded_in_period"], 1)
        self.assertEqual(result["selection"]["abusive_excluded_current"], 1)
        self.assertEqual(result["selection"]["abusive_excluded_previous"], 1)
        self.assertTrue(result["excluded"].empty)
        self.assertNotIn("辱骂", {module["label"] for module in result["modules"]})

    def test_old_cached_suggestion_decisions_are_reprocessed_after_policy_change(self):
        rows, _ = self.selected()
        rows["反馈标签"] = "建议"
        with patch("feedback.ai_client.request_json", side_effect=fake_request):
            process_feedback(rows, self.config, self.root / "processed.csv", lambda **v: None)
        changed = {**self.config, "suggestion_labels": ["产品建议"]}
        with patch("feedback.ai_client.request_json", side_effect=fake_request) as request:
            result = process_feedback(rows, changed, self.root / "processed.csv", lambda **v: None)
        self.assertEqual(request.call_count, 1)
        self.assertTrue(result["是否建议"].eq("否").all())
